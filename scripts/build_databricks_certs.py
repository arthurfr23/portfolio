# Uso: python3 scripts/build_databricks_certs.py <arquivo.xlsx> [snapshot YYYY-MM-DD]
# Requer openpyxl. Gera src/data/databricks-certs-br.json a partir do export do
# Databricks Certified Directory (directory.databrickscertified.com). Aceita também o
# export internacional (com coluna country): nesse caso filtra só o Brasil.
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

from company_names import build_normalizer

OUT = Path(__file__).resolve().parent.parent / 'src' / 'data' / 'databricks-certs-br.json'
ME_PROFILE_ID = '7b018078-2a18-4241-8beb-19a0329ffe50'

CERTS = {
    'Databricks Certified Data Engineer Associate': 'de_assoc',
    'Databricks Certified Data Engineer Professional': 'de_pro',
    'Databricks Certified Data Analyst Associate': 'da_assoc',
    'Databricks Certified Generative AI Engineer Associate': 'genai_assoc',
    'Databricks Certified Machine Learning Engineer Associate': 'ml_assoc',
    'Databricks Certified Machine Learning Professional': 'ml_pro',
    'Databricks Certified Associate Developer for Apache Spark': 'spark',
    'Databricks Certified Associate Developer for Apache Spark 3.0': 'spark',
    'Databricks Certified Associate Developer for Apache Spark 2.4': 'spark',
}

FAMILY_RULES = [
    ('management', r'manager|gerente|\bhead\b|director|diretor|coordenador|coordinator|\bchief\b|\bcto\b|\bceo\b|founder|fundador|\bvp\b|socio|\bjefe\b'),
    ('architecture', r'architect|arquitet|arquitect|solutions? engineer|field engineer|pre.?sales|solutions consultant|specialist solutions'),
    ('analytics_eng', r'analytics engineer|engenheiro de analytics|analytics engineering'),
    ('ai_ml', r'machine learning|\bml\b|mlops|data scien|cientista|cientifico|aprendizaje automatico|\bai\b|\bia\b|genai|\bllm|artificial|inteligencia'),
    ('analyst_bi', r'analyst|analista|\bbi\b|business intelligence'),
    ('consulting', r'consult'),
    ('data_eng', r'data en\w*n\w*er|engenheir[oa] de dados|ingenier\w* de datos|ingenieria de datos|big data|\betl\b|data developer|desenvolvedor de dados|dataops|data platform|engenharia de dados|data engineering|especialista de dados|data specialist|data & analytics|\bdados\b|\bdatos\b|\bdata\b'),
    ('software', r'software|developer|desenvolvedor|programador|backend|full.?stack|engineer|engenheir'),
]

SENIORITY_RULES = [
    ('lead', r'\blead\b|lider|principal|staff|tech lead|especialista|specialist|expert'),
    ('senior', r'senior|\bsr\b|\biii\b|\biv\b'),
    ('mid', r'pleno|\bmid\b|\bii\b|semi.?senior|\bssr\b'),
    ('junior', r'junior|\bjr\b|\bi\b|trainee|estagiari|intern'),
]


def norm(text: str) -> str:
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', text.lower()).strip()


def classify(title: str, rules, default: str) -> str:
    t = norm(title)
    if not t:
        return 'unknown'
    for key, pattern in rules:
        if re.search(pattern, t):
            return key
    return default


def main() -> None:
    src = Path(sys.argv[1])
    snapshot = sys.argv[2] if len(sys.argv) > 2 else date.today().isoformat()
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]

    header = [norm(c.value or '') for c in ws[1]]
    country_col = header.index('country') if 'country' in header else None

    people = []
    unknown_certs = set()
    for row in ws.iter_rows(min_row=2):
        if country_col is not None and (row[country_col].value or '').strip() != 'Brazil':
            continue
        name, count, certs_raw, _, title, workplace = (c.value for c in row[:6])
        link = row[3].hyperlink.target if row[3].hyperlink else None
        if not name or not count or not link:
            continue
        certs = set()
        for line in (certs_raw or '').split('\n'):
            line = line.strip()
            if not line:
                continue
            if line in CERTS:
                certs.add(CERTS[line])
            else:
                unknown_certs.add(line)
        title = (title or '').strip()
        workplace = (workplace or '').strip()
        family = classify(title, FAMILY_RULES, 'other')
        seniority = 'management' if family == 'management' else classify(title, SENIORITY_RULES, 'unspecified')
        person = {
            'n': name.strip(),
            'c': int(count),
            'k': sorted(certs),
            't': title,
            'w': workplace,
            'f': family,
            's': seniority,
            'u': link.rsplit('/', 1)[-1],
        }
        if person['u'] == ME_PROFILE_ID:
            person['me'] = True
        people.append(person)

    companies = build_normalizer([p['w'] for p in people])
    for p in people:
        p['w'] = companies.get(p['w'], p['w'])

    people.sort(key=lambda p: (-p['c'], norm(p['n'])))
    rank = 0
    for i, p in enumerate(people):
        if i == 0 or p['c'] != people[i - 1]['c']:
            rank = i + 1
        p['r'] = rank

    if unknown_certs:
        print('WARN certificações não mapeadas:', sorted(unknown_certs))

    OUT.write_text(json.dumps({'snapshot': snapshot, 'country': 'BR', 'people': people}, ensure_ascii=False, separators=(',', ':')))
    print(f'{len(people)} profissionais -> {OUT}')


if __name__ == '__main__':
    main()
